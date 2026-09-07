# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import copy

TPL = "/Users/yuki/AI/cursor-workspace/slide/slide_templete.pptx"
OUT = "/Users/yuki/AI/pjt_kenpo_web/slides/build/研究の説明と協力のお願い（概要版・事業所担当者向け）.pptx"

NAVY   = RGBColor(0x34,0x6A,0x9D)
PALE   = RGBColor(0xD6,0xE5,0xF2)
INK    = RGBColor(0x1F,0x2A,0x37)
GREY   = RGBColor(0x5B,0x66,0x73)
LGREY  = RGBColor(0x8A,0x94,0xA0)
WHITE  = RGBColor(0xFF,0xFF,0xFF)
BGSOFT = RGBColor(0xEF,0xF2,0xF9)
CORAL  = RGBColor(0xF2,0x6B,0x43)
PINK   = RGBColor(0xFC,0xB6,0xC8)
FONT   = "游ゴシック"
FONTB  = "游ゴシック Medium"

prs = Presentation(TPL)
LAY = {l.name: l for l in prs.slide_master.slide_layouts}

# ---- remove all sample slides -------------------------------------------
xml_slides = prs.slides._sldIdLst
for sld in list(xml_slides):
    rId = sld.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
    prs.part.drop_rel(rId)
    xml_slides.remove(sld)

# ---- helpers -------------------------------------------------------------
def txbox(sl, l, t, w, h, anchor=MSO_ANCHOR.TOP):
    tb = sl.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tb, tf

def setp(p, text, size, color=INK, bold=False, font=None, space_after=0,
         line=None, align=PP_ALIGN.LEFT):
    p.text = text
    p.alignment = align
    if line: p.line_spacing = line
    p.space_after = Pt(space_after)
    for r in p.runs:
        r.font.size = Pt(size); r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = font or (FONTB if bold else FONT)
        # set east-asian typeface too
        rPr = r._r.get_or_add_rPr()
        ea = rPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ea')
        if ea is None:
            ea = rPr.makeelement('{http://schemas.openxmlformats.org/drawingml/2006/main}ea', {})
            rPr.append(ea)
        ea.set('typeface', font or (FONTB if bold else FONT))
    return p

def para(tf, text, size, **kw):
    p = tf.paragraphs[0] if (len(tf.paragraphs) == 1 and not tf.paragraphs[0].text) else tf.add_paragraph()
    return setp(p, text, size, **kw)

def rect(sl, l, t, w, h, fill=None, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=None, lw=1.0):
    s = sl.shapes.add_shape(shape, Inches(l), Inches(t), Inches(w), Inches(h))
    if adj is not None:
        try: s.adjustments[0] = adj
        except Exception: pass
    if fill is None: s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None: s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(lw)
    s.shadow.inherit = False
    return s

SECTIONS = ["研究について", "ご協力の内容", "安心のために", "お手続き"]

def add_slide(title, sub, section_idx, page):
    sl = prs.slides.add_slide(LAY["タイトルとコンテンツ"])
    # fill title placeholder, drop body placeholder
    for ph in list(sl.placeholders):
        if ph.placeholder_format.idx == 0:
            tf = ph.text_frame; tf.word_wrap = True
            setp(tf.paragraphs[0], title, 28, color=INK, bold=True)
            ph.top, ph.left = Inches(0.30), Inches(0.62)
            ph.width, ph.height = Inches(9.6), Inches(0.62)
        else:
            ph._element.getparent().remove(ph._element)
    if sub:
        _, tf = txbox(sl, 0.64, 0.95, 10.0, 0.35)
        para(tf, sub, 13.5, color=GREY)
    # section tabs (right edge) — active = navy
    for i in range(4):
        s = rect(sl, 12.90, 1.00 + i*0.39, 0.30, 0.56,
                 fill=(NAVY if i == section_idx else PALE),
                 shape=MSO_SHAPE.ROUND_2_SAME_RECTANGLE)
        s.rotation = 90
    _, tf = txbox(sl, 12.84, 2.68, 0.44, 2.06)
    para(tf, SECTIONS[section_idx], 11, color=NAVY, bold=True)
    # page number
    _, tf = txbox(sl, 11.6, 7.02, 1.1, 0.3)
    para(tf, f"{page} / 7", 10, color=LGREY, align=PP_ALIGN.RIGHT)
    return sl

def card(sl, l, t, w, h, fill=BGSOFT):
    return rect(sl, l, t, w, h, fill=fill, adj=0.045)

def numbadge(sl, l, t, n, d=0.44, color=NAVY):
    s = rect(sl, l, t, d, d, fill=color, shape=MSO_SHAPE.OVAL)
    tf = s.text_frame; tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    setp(tf.paragraphs[0], str(n), 15, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    return s

# =========================================================================
# 1. 表紙
# =========================================================================
sl = prs.slides.add_slide(LAY["1_ユーザー設定レイアウト"])
for ph in list(sl.placeholders):
    idx = ph.placeholder_format.idx
    if idx == 0:
        tf = ph.text_frame; tf.word_wrap = True
        setp(tf.paragraphs[0], "研究の説明と協力のお願い", 40, color=INK, bold=True)
        ph.left, ph.top, ph.width, ph.height = Inches(2.13), Inches(3.25), Inches(9.07), Inches(1.0)
    elif idx == 13:
        tf = ph.text_frame; tf.word_wrap = True
        setp(tf.paragraphs[0], "事業所のご担当者の皆さまへ　＜概要版＞", 17, color=NAVY, bold=True)
    elif idx == 12:
        tf = ph.text_frame; tf.word_wrap = True
        setp(tf.paragraphs[0], "横浜市立大学医学部 公衆衛生学教室　研究責任者 / 荒川 裕貴", 12, color=GREY)
        ph.left, ph.top, ph.width, ph.height = Inches(6.05), Inches(6.42), Inches(6.6), Inches(0.34)
        tf.word_wrap = False
        ph.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT

_, tf = txbox(sl, 2.16, 4.32, 9.03, 0.4)
para(tf, "企業の構造的・文化的要因と働く男女の健康の関連に関する疫学研究", 14, color=GREY)

# partner strip
y = 5.15
rect(sl, 2.13, y, 9.07, 0.9, fill=WHITE, adj=0.10)
rect(sl, 2.13, y, 0.055, 0.9, fill=NAVY, shape=MSO_SHAPE.RECTANGLE)
cols = [("研究代表機関", "横浜市立大学医学部 公衆衛生学教室"),
        ("連携機関", "横浜市 健康福祉局健康推進課 ／ 協会けんぽ神奈川支部")]
xs = [2.50, 6.35]
for (h, b), x in zip(cols, xs):
    _, tf = txbox(sl, x, y + 0.19, 3.6, 0.24)
    para(tf, h, 10, color=NAVY, bold=True)
    _, tf = txbox(sl, x, y + 0.46, 4.6, 0.30)
    para(tf, b, 11.5, color=INK)

# =========================================================================
# 2. なぜこの研究を行うのか
# =========================================================================
sl = add_slide("なぜ、この研究を行うのか", "働く人の健康と、その人が働く会社の環境との関係を明らかにします", 0, 2)

card(sl, 0.62, 1.55, 5.85, 3.72)
rect(sl, 0.62, 1.55, 5.85, 0.075, fill=CORAL, shape=MSO_SHAPE.RECTANGLE)
_, tf = txbox(sl, 1.00, 1.95, 5.1, 0.35)
para(tf, "見過ごされがちな健康課題", 17, color=INK, bold=True)
_, tf = txbox(sl, 1.00, 2.48, 5.1, 0.5)
para(tf, "仕事や生活に大きな影響を与えることが知られています。", 12, color=GREY, line=1.5)
items = ["月経困難症・更年期症状などの体の不調", "うつなどの心の不調", "職場での孤独感・孤立感"]
yy = 3.20
for it in items:
    rect(sl, 1.02, yy + 0.115, 0.10, 0.10, fill=CORAL, shape=MSO_SHAPE.OVAL)
    _, tf = txbox(sl, 1.32, yy, 4.8, 0.34)
    para(tf, it, 13, color=INK)
    yy += 0.62

card(sl, 6.75, 1.55, 5.85, 3.72)
rect(sl, 6.75, 1.55, 5.85, 0.075, fill=NAVY, shape=MSO_SHAPE.RECTANGLE)
_, tf = txbox(sl, 7.13, 1.95, 5.1, 0.35)
para(tf, "まだ分かっていないこと", 17, color=INK, bold=True)
_, tf = txbox(sl, 7.13, 2.48, 5.1, 1.6)
para(tf, "会社の制度や体制、職場の雰囲気といった「働く環境」が、"
         "一人ひとりの健康とどう結びついているのか。", 13, color=INK, line=1.6)
rect(sl, 7.13, 3.72, 5.1, 1.10, fill=WHITE, adj=0.08)
_, tf = txbox(sl, 7.42, 3.95, 4.55, 0.66)
para(tf, "その関係は、まだ科学的に十分示されていません。", 13, color=NAVY, bold=True, line=1.4)

_, tf = txbox(sl, 0.64, 5.60, 11.9, 0.4)
para(tf, "だからこそ、事業所と働く方の双方からデータを集め、両者のつながりを検証します。", 12.5, color=GREY)

# =========================================================================
# 3. 研究の目的と意義
# =========================================================================
sl = add_slide("研究の目的と意義", "データを用いて、働く環境と健康のつながりを明らかにすることを目指します", 0, 3)
goals = [("1", "関連を明らかにする", "就労先企業の機能的・構造的・文化的な要因が、心・体・社会的な健康にどう影響するかをデータで解明します。"),
         ("2", "根拠をつくる", "企業や保険者が健康づくりの取り組みを検討・改善する際に使える、科学的な根拠を示します。"),
         ("3", "健康づくりに活かす", "働く男女すべての健康づくりに役立つ、具体的な手がかりにつなげます。")]
x = 0.62
for n, h, b in goals:
    card(sl, x, 1.62, 3.86, 2.72)
    numbadge(sl, x + 0.42, 2.05, n)
    _, tf = txbox(sl, x + 0.42, 2.78, 3.1, 0.35)
    para(tf, h, 16, color=INK, bold=True)
    _, tf = txbox(sl, x + 0.42, 3.28, 3.05, 1.8)
    para(tf, b, 12, color=GREY, line=1.55)
    x += 4.13

rect(sl, 0.62, 4.86, 11.98, 0.82, fill=PALE, adj=0.10)
_, tf = txbox(sl, 1.00, 5.09, 11.3, 0.4)
para(tf, "研究成果は、企業・保険者が働く人の健康づくりを進めるための判断材料として活用されることが期待されます。",
     12.5, color=NAVY, bold=True)

# =========================================================================
# 4. 対象と調査の流れ
# =========================================================================
sl = add_slide("ご協力いただく方と、調査の流れ", "オンライン質問票（ウェブアンケート）にご回答いただきます", 1, 4)

card(sl, 0.62, 1.55, 11.98, 1.62, fill=BGSOFT)
rect(sl, 0.62, 1.55, 0.075, 1.62, fill=NAVY, shape=MSO_SHAPE.RECTANGLE)
_, tf = txbox(sl, 1.02, 1.80, 3.0, 0.3)
para(tf, "対象となる事業所", 15, color=INK, bold=True)
_, tf = txbox(sl, 1.02, 2.20, 5.2, 0.7)
para(tf, "協会けんぽ神奈川支部に加入し、説明会等を通じて参加を希望した、"
         "日本に所在地を持つ事業所", 12.5, color=GREY, line=1.5)
rect(sl, 6.60, 1.83, 5.72, 1.08, fill=WHITE, adj=0.09)
_, tf = txbox(sl, 6.92, 2.02, 5.1, 0.3)
para(tf, "ご回答は、人事・労務・健康づくりのご担当者にお願いします", 12.5, color=NAVY, bold=True)
_, tf = txbox(sl, 6.92, 2.38, 5.1, 0.3)
para(tf, "あわせて、そこで働く18〜74歳の方にもご協力をお願いしています", 11.5, color=GREY)

_, tf = txbox(sl, 0.64, 3.42, 6.0, 0.3)
para(tf, "ご回答までの流れ", 15, color=INK, bold=True)

steps = [("1", "ご案内", "研究用ホームページで、\n本説明事項をご確認ください。"),
         ("2", "同意", "内容にご納得いただけたら、\n画面上で同意の手続きを行います。"),
         ("3", "ご回答", "オンライン質問票にご回答。\n所要時間は約20分です。"),
         ("4", "結果のお返し", "全体の集計結果と自社の回答\nとの比較をお返しします。")]
x = 0.62
for i, (n, h, b) in enumerate(steps):
    card(sl, x, 3.88, 2.78, 2.05, fill=BGSOFT)
    numbadge(sl, x + 0.30, 4.16, n, d=0.38)
    _, tf = txbox(sl, x + 0.80, 4.22, 1.9, 0.3)
    para(tf, h, 13.5, color=INK, bold=True)
    _, tf = txbox(sl, x + 0.30, 4.78, 2.25, 1.0)
    tfp = None
    for j, ln in enumerate(b.split("\n")):
        para(tf, ln, 11.5, color=GREY, line=1.45)
    if i < 3:
        a, atf = txbox(sl, x + 2.83, 4.72, 0.32, 0.3)
        para(atf, "▶", 11, color=PALE, align=PP_ALIGN.CENTER)
    x += 3.06

_, tf = txbox(sl, 0.64, 6.18, 11.9, 0.35)
para(tf, "血液などの生体試料は扱いません。ご回答は識別コードを付けた研究用データとして取り扱います。", 12, color=GREY)

# =========================================================================
# 5. おうかがいする内容
# =========================================================================
sl = add_slide("おうかがいする内容", "所要時間の目安は約20分です（設問数により変動します）", 1, 5)

card(sl, 0.62, 1.58, 7.15, 4.05)
rect(sl, 0.62, 1.58, 7.15, 0.075, fill=NAVY, shape=MSO_SHAPE.RECTANGLE)
_, tf = txbox(sl, 1.02, 1.98, 6.4, 0.32)
para(tf, "事業所のご担当者へ", 17, color=INK, bold=True)
_, tf = txbox(sl, 1.02, 2.40, 6.4, 0.3)
para(tf, "貴事業所の体制・制度についてうかがいます", 12, color=GREY)
qs = ["業種・従業員数などの基本情報", "健康づくりの体制",
      "休暇制度の内容と利用状況", "職員交流の取り組み", "性差に配慮した取り組み"]
yy = 2.92
for q in qs:
    rect(sl, 1.04, yy + 0.115, 0.10, 0.10, fill=NAVY, shape=MSO_SHAPE.OVAL)
    _, tf = txbox(sl, 1.34, yy, 6.0, 0.32)
    para(tf, q, 13, color=INK)
    yy += 0.52

card(sl, 8.05, 1.58, 4.55, 4.05, fill=RGBColor(0xF7,0xF9,0xFC))
rect(sl, 8.05, 1.58, 4.55, 0.075, fill=PINK, shape=MSO_SHAPE.RECTANGLE)
_, tf = txbox(sl, 8.42, 1.98, 3.9, 0.32)
para(tf, "働く方へ", 17, color=INK, bold=True)
_, tf = txbox(sl, 8.42, 2.40, 3.9, 0.3)
para(tf, "従業員の皆さまにうかがう内容", 12, color=GREY)
qs2 = ["年齢・性別・勤務形態などの背景", "生活習慣", "メンタルヘルス",
       "孤独感、労働生産性", "女性には月経困難症・更年期症状"]
yy = 2.92
for q in qs2:
    rect(sl, 8.44, yy + 0.115, 0.10, 0.10, fill=PINK, shape=MSO_SHAPE.OVAL)
    _, tf = txbox(sl, 8.74, yy, 3.7, 0.32)
    para(tf, q, 12.5, color=INK)
    yy += 0.52

rect(sl, 0.62, 5.90, 11.98, 0.72, fill=PALE, adj=0.11)
_, tf = txbox(sl, 1.00, 6.09, 11.3, 0.35)
para(tf, "答えたくない質問には、回答いただかなくて構いません。ご回答は自由意思に基づくものです。",
     12.5, color=NAVY, bold=True)

# =========================================================================
# 6. 利益・負担 と 個人情報 と 任意性
# =========================================================================
sl = add_slide("ご参加による利益・ご負担と、個人情報の取り扱い", "利益と負担の両方をお伝えし、情報は厳重に管理します", 2, 6)

# left: merits
card(sl, 0.62, 1.55, 3.86, 2.85)
_, tf = txbox(sl, 0.98, 1.82, 3.2, 0.3)
para(tf, "期待される利益", 15, color=INK, bold=True)
for i, t in enumerate(["自社・ご自身の状況を振り返る機会になります",
                       "全体の集計値との比較をお返しします",
                       "健康づくりに役立つ情報提供動画をご案内します"]):
    rect(sl, 1.00, 2.32 + i*0.66 + 0.10, 0.09, 0.09, fill=NAVY, shape=MSO_SHAPE.OVAL)
    _, tf = txbox(sl, 1.28, 2.32 + i*0.66, 2.95, 0.6)
    para(tf, t, 11.5, color=GREY, line=1.4)

# middle: burden
card(sl, 4.68, 1.55, 3.86, 2.85)
_, tf = txbox(sl, 5.04, 1.82, 3.2, 0.3)
para(tf, "予測されるご負担", 15, color=INK, bold=True)
for i, t in enumerate(["ご回答に20分ほどのお時間をいただきます",
                       "月経・更年期・メンタルヘルスに関する質問を含みます",
                       "身体的な侵襲はなく、費用の負担もありません"]):
    rect(sl, 5.06, 2.32 + i*0.66 + 0.10, 0.09, 0.09, fill=CORAL, shape=MSO_SHAPE.OVAL)
    _, tf = txbox(sl, 5.34, 2.32 + i*0.66, 2.95, 0.6)
    para(tf, t, 11.5, color=GREY, line=1.4)

# right: privacy
card(sl, 8.74, 1.55, 3.86, 2.85, fill=PALE)
_, tf = txbox(sl, 9.10, 1.82, 3.2, 0.3)
para(tf, "個人情報の取り扱い", 15, color=NAVY, bold=True)
for i, t in enumerate(["識別コードを付けて管理し、氏名等とは分けて保管します",
                       "対応表を研究機関の外へ出すことはありません",
                       "事業者・協会けんぽ・横浜市に個人情報は渡りません"]):
    rect(sl, 9.12, 2.32 + i*0.66 + 0.10, 0.09, 0.09, fill=NAVY, shape=MSO_SHAPE.OVAL)
    _, tf = txbox(sl, 9.40, 2.32 + i*0.66, 2.95, 0.6)
    para(tf, t, 11.5, color=INK, bold=False, line=1.4)

# voluntary band
rect(sl, 0.62, 4.62, 11.98, 1.95, fill=WHITE, line=NAVY, adj=0.06, lw=1.25)
rect(sl, 0.62, 4.62, 0.075, 1.95, fill=NAVY, shape=MSO_SHAPE.RECTANGLE)
_, tf = txbox(sl, 1.02, 4.88, 8.0, 0.35)
para(tf, "ご参加は自由意思です", 18, color=NAVY, bold=True)
_, tf = txbox(sl, 1.02, 5.32, 11.2, 0.35)
para(tf, "参加を断っても、途中でやめても、不利益を受けることはありません。", 13, color=INK)
vs = [("同意はいつでも撤回できます", "同意後でも、いつでも参加をやめられます。"),
      ("不利益はありません", "断っても撤回しても、不利益は生じません。"),
      ("撤回のご連絡先", "次ページのお問い合わせ先までご連絡ください。")]
x = 1.02
for h, b in vs:
    _, tf = txbox(sl, x, 5.80, 3.6, 0.28)
    para(tf, "・" + h, 12, color=NAVY, bold=True)
    _, tf = txbox(sl, x + 0.19, 6.09, 3.5, 0.28)
    para(tf, b, 11, color=GREY)
    x += 3.78

# =========================================================================
# 7. お手続きとお問い合わせ
# =========================================================================
sl = add_slide("ご参加の手続きとお問い合わせ先", "内容をよくご理解いただいたうえで、ご自身の意思でご判断ください", 3, 7)

_, tf = txbox(sl, 0.64, 1.52, 6.0, 0.3)
para(tf, "ご参加いただける場合の手順", 15, color=INK, bold=True)
proc = [("1", "元の画面に戻り、「上の項目を十分理解したうえで、研究への協力に同意します」にチェックを入れます。"),
        ("2", "画面を進み、必要事項を入力します。\n事業所担当者の方：氏名、役職、メールアドレス"),
        ("3", "オンライン質問票にご回答・送信いただくと、本研究への参加に\n同意いただいたものとみなします（電磁的方法による同意）。")]
yy = 1.95
for n, b in proc:
    card(sl, 0.62, yy, 7.15, 0.94, fill=BGSOFT)
    numbadge(sl, 0.92, yy + 0.26, n, d=0.40)
    _, tf = txbox(sl, 1.52, yy + 0.20, 6.05, 0.62)
    for ln in b.split("\n"):
        para(tf, ln, 12, color=INK, line=1.35)
    yy += 1.06

# contact card
rect(sl, 8.05, 1.95, 4.55, 3.12, fill=NAVY, adj=0.07)
_, tf = txbox(sl, 8.45, 2.22, 3.8, 0.3)
para(tf, "お問い合わせ先", 15, color=WHITE, bold=True)
rect(sl, 8.45, 2.62, 0.55, 0.035, fill=PINK, shape=MSO_SHAPE.RECTANGLE)
info = [("機関名", "横浜市立大学医学部 公衆衛生学教室"),
        ("研究責任者", "荒川 裕貴"),
        ("電話", "045-787-2610"),
        ("メール", "arakawa.yuk.tq@yokohama-cu.ac.jp")]
yy = 2.85
for h, b in info:
    _, tf = txbox(sl, 8.45, yy, 3.9, 0.22)
    para(tf, h, 9.5, color=PALE)
    _, tf = txbox(sl, 8.45, yy + 0.23, 3.95, 0.28)
    para(tf, b, 12, color=WHITE, bold=True)
    yy += 0.56

# fine print / ethics footer
rect(sl, 0.62, 5.30, 11.98, 1.28, fill=RGBColor(0xF7,0xF9,0xFC), adj=0.07)
_, tf = txbox(sl, 1.00, 5.50, 11.3, 0.25)
para(tf, "そのほかの重要事項", 11.5, color=NAVY, bold=True)
fine = ("本研究は横浜市立大学「人を対象とする生命科学・医学系研究倫理委員会」の審査・承認と研究機関の長の許可を得て実施します。"
        "研究期間は実施許可日〜2031年3月31日（登録期間は2028年3月31日まで）。"
        "情報は研究終了報告日から5年、または最終公表報告日から3年のいずれか遅い日まで保管し、復元できない方法で廃棄します。")
fine2 = ("将来、他の研究への利用や研究機関への提供を行う場合は、改めて倫理審査委員会の承認と研究機関の長の許可を得たうえで適切に手続します。"
         "日本学術振興会 科学研究費助成事業（若手研究）および横浜市受託研究により実施し、利益相反はありません。"
         "成果は学会・論文等で公表しますが、個人が特定される情報は公表しません。")
_, tf = txbox(sl, 1.00, 5.78, 11.25, 0.72)
para(tf, fine, 9.5, color=GREY, line=1.45)
para(tf, fine2, 9.5, color=GREY, line=1.45)

prs.save(OUT)
print("SAVED:", OUT, "slides:", len(prs.slides.__iter__.__self__._sldIdLst))
