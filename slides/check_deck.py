"""生成した pptx の構造検査。

skill 付属の validate.py は Python 3.10 以上を要求するため、同等の検査を
この環境（3.9）で動く形で行う。
  - ZIP 健全性と XML の整形式性
  - ppt/slides/*.xml の XSD 検証
  - 図形がスライド外へはみ出していないか
  - テキストが枠に収まるかの概算
"""

import sys
import zipfile

from lxml import etree
from pptx import Presentation
from pptx.util import Emu

SKILL = "/Users/yuki/AI/.agents/skills/pptx/scripts/office"
SCHEMA = f"{SKILL}/schemas/ISO-IEC29500-4_2016/pml.xsd"

path = sys.argv[1]
problems = []
notes = []

# ---- ZIP と XML ----
z = zipfile.ZipFile(path)
bad = z.testzip()
if bad:
    problems.append(f"ZIP破損: {bad}")
for name in z.namelist():
    if name.endswith((".xml", ".rels")):
        try:
            etree.fromstring(z.read(name))
        except etree.XMLSyntaxError as e:
            problems.append(f"XML不正 {name}: {e}")

# ---- XSD ----
schema = etree.XMLSchema(etree.parse(SCHEMA))
# 同梱スキーマは ISO strict。pptxgenjs は箇条書きに transitional 形式の
# buSzPct val="100000"（1/1000%単位）を書き、PowerPoint はこれを読む。
KNOWN_STRICT_ONLY = "buSzPct"
for name in sorted(n for n in z.namelist() if n.startswith("ppt/slides/slide")):
    doc = etree.fromstring(z.read(name))
    if not schema.validate(doc.getroottree()):
        for e in schema.error_log:
            if KNOWN_STRICT_ONLY in e.message:
                continue
            problems.append(f"XSD {name}: {e.message}")

# ---- 図形の配置とテキスト量 ----
prs = Presentation(path)
SW, SH = prs.slide_width, prs.slide_height
EMU_IN = 914400.0


def wide(ch):
    """全角ならおよそ1文字幅、半角なら約0.55文字幅として数える。"""
    return 0.55 if ord(ch) < 0x2E80 else 1.0


for idx, slide in enumerate(prs.slides, 1):
    for sh in slide.shapes:
        if sh.left is None:
            continue
        l, t = sh.left, sh.top
        r, b = l + (sh.width or 0), t + (sh.height or 0)
        if l < 0 or t < 0 or r > SW or b > SH:
            problems.append(
                f"S{idx} スライド外: {sh.shape_type} "
                f"({l/EMU_IN:.2f},{t/EMU_IN:.2f})-({r/EMU_IN:.2f},{b/EMU_IN:.2f})"
            )
        margin = Emu(int(0.5 * EMU_IN))
        if l < margin or t < margin or SW - r < margin or SH - b < margin:
            if not (sh.width == SW and sh.height == SH):
                notes.append(
                    f"S{idx} 余白0.5\"未満: ({l/EMU_IN:.2f},{t/EMU_IN:.2f})-"
                    f"({r/EMU_IN:.2f},{b/EMU_IN:.2f})"
                )

        if not sh.has_text_frame:
            continue
        tf = sh.text_frame
        text = tf.text
        if not text.strip():
            continue
        size = None
        for p in tf.paragraphs:
            for run in p.runs:
                if run.font.size:
                    size = run.font.size.pt
                    break
            if size:
                break
        if not size:
            continue
        box_w_in = (sh.width or 0) / EMU_IN
        box_h_in = (sh.height or 0) / EMU_IN
        char_in = size / 72.0
        per_line = max(1, int(box_w_in / char_in))
        lines = 0
        for para in text.split("\n"):
            units = sum(wide(c) for c in para)
            lines += max(1, -(-int(units + 0.999) // per_line))
        need_in = lines * (size * 1.62) / 72.0
        if need_in > box_h_in + 0.04:
            problems.append(
                f"S{idx} テキスト超過の可能性: {size}pt, 推定{lines}行 "
                f"必要{need_in:.2f}\" > 枠{box_h_in:.2f}\" :: {text[:34]}…"
            )

print(f"slides: {len(prs.slides.__iter__.__self__._sldIdLst)}")
if notes:
    print("\n[margin notes]")
    for n in sorted(set(notes)):
        print(" -", n)
print("\n[problems]" if problems else "\nOK: 問題は検出されませんでした")
for p in problems:
    print(" -", p)
sys.exit(1 if problems else 0)
